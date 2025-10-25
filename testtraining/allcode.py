# simple login
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
import time


service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service)

driver.get("https://practicetestautomation.com/practice-test-login/")
driver.maximize_window()
time.sleep(2)  

username_input = driver.find_element(By.XPATH, "//input[@id='username']")
username_input.send_keys("student")
time.sleep(1)

password_input = driver.find_element(By.XPATH, "//input[@id='password']")
password_input.send_keys("Password123")
time.sleep(1)

submit_button = driver.find_element(By.XPATH, "//button[@id='submit']")
submit_button.click()
time.sleep(3)
driver.quit()

#loop save data excel  4 atteppmt---------------------------
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
from openpyxl import Workbook, load_workbook
from datetime import datetime
import os
import time

# Step 1: Auto-create Excel for input
input_file = "login_input.xlsx"
if not os.path.exists(input_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Credentials"
    ws.append(["Username", "Password"])
    ws.append(["wronguser", "Password123"])      # wrong username
    ws.append(["student", "wrongpass"])          # wrong password
    ws.append(["wronguser", "wrongpass"])        # both wrong
    ws.append(["student", "Password123"])        # correct
    wb.save(input_file)
    print(f"Input Excel created: {input_file}")

# Step 2: Load input Excel
input_wb = load_workbook(input_file)
input_ws = input_wb.active

# Step 3: Prepare result Excel
result_file = "login_results.xlsx"
wb_result = Workbook()
ws_result = wb_result.active
ws_result.title = "Login Results"
ws_result.append(["Timestamp", "Username", "Password", "Result", "Screenshot Path"])

# Step 4: Setup Selenium
service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service)
driver.maximize_window()
login_url = "https://practicetestautomation.com/practice-test-login/"

# Step 5: Loop over Excel rows
for row in input_ws.iter_rows(min_row=2, values_only=True):
    username_val = row[0]
    password_val = row[1]

    driver.get(login_url)
    time.sleep(1)

    # Fill form
    driver.find_element(By.ID, "username").clear()
    driver.find_element(By.ID, "username").send_keys(username_val)
    driver.find_element(By.ID, "password").clear()
    driver.find_element(By.ID, "password").send_keys(password_val)
    driver.find_element(By.ID, "submit").click()
    time.sleep(2)

    # Step 6: Check result
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    screenshot_name = f"screenshot_{username_val}_{timestamp}.png"
    screenshot_path = os.path.join(os.getcwd(), screenshot_name)
    driver.save_screenshot(screenshot_path)

    current_url = driver.current_url
    if "practicetestautomation.com/logged-in-successfully/" in current_url:
        result = "SUCCESS"
    else:
        # Get error text if visible
        try:
            error_msg = driver.find_element(By.ID, "error").text
            result = f"FAILED - {error_msg}"
        except:
            result = "FAILED - Unknown Error"

    print(f"{username_val}/{password_val} -> {result}")
    ws_result.append([timestamp, username_val, password_val, result, screenshot_path])

# Step 7: Save results
wb_result.save(result_file)
print(f"Results saved in {result_file}")

driver.quit()
#code data from direct excel manual data  auto screen  result seprtae excel sheet final login
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
from openpyxl import load_workbook, Workbook
from datetime import datetime
import os
import time

# Excel file paths
input_file = "login_input.xlsx"
result_file = "login_results.xlsx"

# Step 1: Check input Excel exists
if not os.path.exists(input_file):
    print(f"Input file '{input_file}' not found. Please create it with Username and Password columns.")
    exit()

# Step 2: Load input Excel
input_wb = load_workbook(input_file)
input_ws = input_wb.active

# Step 3: Prepare result Excel
wb_result = Workbook()
ws_result = wb_result.active
ws_result.title = "Login Results"
ws_result.append(["Timestamp", "Username", "Password", "Result", "Screenshot Path"])

# Step 4: Setup Selenium
service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service)
driver.maximize_window()
login_url = "https://practicetestautomation.com/practice-test-login/"

# Step 5: Loop over Excel rows
for row in input_ws.iter_rows(min_row=2, values_only=True):
    username_val = row[0]
    password_val = row[1]

    driver.get(login_url)
    time.sleep(1)

    # Fill form
    driver.find_element(By.ID, "username").clear()
    driver.find_element(By.ID, "username").send_keys(username_val)
    driver.find_element(By.ID, "password").clear()
    driver.find_element(By.ID, "password").send_keys(password_val)
    driver.find_element(By.ID, "submit").click()
    time.sleep(2)

    # Step 6: Check result
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    screenshot_name = f"screenshot_{username_val}_{timestamp}.png"
    screenshot_path = os.path.join(os.getcwd(), screenshot_name)
    driver.save_screenshot(screenshot_path)

    current_url = driver.current_url
    if "practicetestautomation.com/logged-in-successfully/" in current_url:
        result = "SUCCESS"
    else:
        # Get error text if visible
        try:
            error_msg = driver.find_element(By.ID, "error").text
            result = f"FAILED - {error_msg}"
        except:
            result = "FAILED - Unknown Error"

    print(f"{username_val}/{password_val} -> {result}")
    ws_result.append([timestamp, username_val, password_val, result, screenshot_path])

# Step 7: Save results
wb_result.save(result_file)
print(f"Results saved in {result_file}")

driver.quit()
# navigation  automation---------------------------------------------------------------------------------
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
import time

# Step 1: Setup Selenium
service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service)
driver.maximize_window()

# Step 2: Navigate to OrangeHRM demo login
url = "https://opensource-demo.orangehrmlive.com/web/index.php/auth/login"
driver.get(url)
time.sleep(2)

# Step 3: Click element using your XPath
xpath = '//*[@id="app"]/div[1]/div/div[1]/div/div[2]/div[2]/form/div[4]/p'
element = driver.find_element(By.XPATH, xpath)
element.click()
print("Element clicked!")
time.sleep(2)

# Step 4: Browser navigation examples

# Go back
driver.back()
print("Went back")
time.sleep(2)

# Go forward
driver.forward()
print("Went forward")
time.sleep(2)

# Refresh page
driver.refresh()
print("Page refreshed")
time.sleep(2)

# Step 5: Close browser
driver.quit()

# checkbox
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager
import time

# Setup Selenium
service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service)
driver.maximize_window()

driver.get("https://www.tutorialspoint.com/selenium/practice/selenium_automation_practice.php")

# Wait until the checkbox is clickable
wait = WebDriverWait(driver, 10)
checkbox = wait.until(EC.element_to_be_clickable((By.XPATH, '//*[@id="hobbies"]')))
checkbox.click()
print("Checkbox clicked using EC wait!")

time.sleep(2)
driver.quit()

#alert popup--------------------------------------------------------------------------------------------------
import time
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

# -------------------------
# Inicializar driver
# -------------------------
driver = webdriver.Chrome()
driver.maximize_window()

url = "https://demo.automationtesting.in/Alerts.html"
driver.get(url)
time.sleep(2)

# -------------------------
# Crear WebDriverWait
# -------------------------
wait = WebDriverWait(driver, 10)

# -------------------------
# Paso 1: Simple Alert
# -------------------------
simple_alert = wait.until(
    EC.element_to_be_clickable((By.XPATH, '//*[@id="OKTab"]/button'))
)
simple_alert.click()
time.sleep(1)

alertas = driver.switch_to.alert
print("Paso 1: OK alert")
alertas.accept()
time.sleep(2)

# -------------------------
# Paso 2: Confirmation Alert
# -------------------------
# Click en la pestaña "Alert with OK & Cancel"
driver.find_element(By.XPATH, '/html/body/div[1]/div/div/div/div[1]/ul/li[2]/a').click()

# Esperar a que el botón sea clickeable
confirm_alert = wait.until(
    EC.element_to_be_clickable((By.XPATH, '//*[@id="CancelTab"]/button'))
)
confirm_alert.click()
time.sleep(1)

# Aceptar confirm alert
confirm_alerts = driver.switch_to.alert
print("Paso 2: Confirmation alert - Accept")
confirm_alerts.accept()
time.sleep(2)

# Volver a hacer clic para probar dismiss
confirm_alert.click()
confirm_alerts = driver.switch_to.alert
print("Paso 2: Confirmation alert - Dismiss")
confirm_alerts.dismiss()
time.sleep(2)

driver.quit()


# mouse hover------------------------------------------------------------------
import time
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

driver = webdriver.Chrome()
driver.maximize_window()

url = "https://practice.expandtesting.com/hovers"
driver.get(url)

wait = WebDriverWait(driver, 10)
actions = ActionChains(driver)

hover_element1 = wait.until(EC.visibility_of_element_located(
    (By.XPATH, '//*[@id="core"]/div/div/div[1]/img')
))
driver.execute_script("arguments[0].scrollIntoView(true);", hover_element1)
actions.move_to_element(hover_element1).perform()
print("Paso 1: OK")
time.sleep(2)

hover_element2 = wait.until(EC.visibility_of_element_located(
    (By.XPATH, '//*[@id="core"]/div/div/div[2]/img')
))
driver.execute_script("arguments[0].scrollIntoView(true);", hover_element2)
actions.move_to_element(hover_element2).perform()
print("Paso 2: OK")
time.sleep(2)

hover_element3 = wait.until(EC.visibility_of_element_located(
    (By.XPATH, '//*[@id="core"]/div/div/div[3]/img')
))
driver.execute_script("arguments[0].scrollIntoView(true);", hover_element3)
actions.move_to_element(hover_element3).perform()
print("Paso 3: OK")
time.sleep(2)


driver.quit()
