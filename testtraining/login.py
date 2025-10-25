from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
from openpyxl import load_workbook, Workbook
from datetime import datetime
import os
import time

# Excel file paths
input_file = "login_input.xlsx"
result_file = "login_results.xlsx"

# Step 1: Check input Excel exists
if not os.path.exists(input_file):
    print(f"Input file '{input_file}' not found. Please create it with Username and Password columns.")
    exit()

# Step 2: Load input Excel
input_wb = load_workbook(input_file)
input_ws = input_wb.active

# Step 3: Prepare result Excel
wb_result = Workbook()
ws_result = wb_result.active
ws_result.title = "Login Results"
ws_result.append(["Timestamp", "Username", "Password", "Result", "Screenshot Path"])

# Step 4: Setup Selenium
service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service)
driver.maximize_window()
login_url = "https://practicetestautomation.com/practice-test-login/"

# Step 5: Loop over Excel rows
for row in input_ws.iter_rows(min_row=2, values_only=True):
    username_val = row[0]
    password_val = row[1]

    driver.get(login_url)
    time.sleep(1)

    # Fill form
    driver.find_element(By.ID, "username").clear()
    driver.find_element(By.ID, "username").send_keys(username_val)
    driver.find_element(By.ID, "password").clear()
    driver.find_element(By.ID, "password").send_keys(password_val)
    driver.find_element(By.ID, "submit").click()
    time.sleep(2)

    # Step 6: Check result
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    screenshot_name = f"screenshot_{username_val}_{timestamp}.png"
    screenshot_path = os.path.join(os.getcwd(), screenshot_name)
    driver.save_screenshot(screenshot_path)

    current_url = driver.current_url
    if "practicetestautomation.com/logged-in-successfully/" in current_url:
        result = "SUCCESS"
    else:
        # Get error text if visible
        try:
            error_msg = driver.find_element(By.ID, "error").text
            result = f"FAILED - {error_msg}"
        except:
            result = "FAILED - Unknown Error"

    print(f"{username_val}/{password_val} -> {result}")
    ws_result.append([timestamp, username_val, password_val, result, screenshot_path])

# Step 7: Save results
wb_result.save(result_file)
print(f"Results saved in {result_file}")

driver.quit()
